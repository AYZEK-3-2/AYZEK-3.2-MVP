# -*- coding: utf-8 -*-
"""
Çıktı dosyaları üreticisi.

Üç xlsx dosyasını yarışma örnek formatına birebir uygun şekilde yazar:
  - Tahminlenen_Talep.xlsx
  - Araç_Planlama.xlsx   (toplam maliyet satırı dahil)
  - Karşılanamayan_Yük.xlsx

Kullanım:
    from src.output_writer import write_all
    toplam = write_all(df_tahmin_out, tam_plan, unfulfilled, OUTPUT_FILES)
"""
import pandas as pd
from pathlib import Path
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

# Yarışma formatına göre kesin kolon sırası
TAHMIN_KOLONLAR = ["Tarih", "Çıkış TM", "Varış TM", "Tahmin Edilen Desi"]
PLAN_KOLONLAR   = ["Tarih", "Araç Tipi", "Çıkış TM", "Varış TM", "Güzergah", "Atanan Desi", "Maliyet"]
UNFUL_KOLONLAR  = ["Tarih", "Çıkış TM", "Varış TM", "Karşılanamayan Desi", "Sebep"]

_HEADER_FILL = PatternFill(fill_type="solid", fgColor="D9E1F2")
_HEADER_FONT = Font(bold=True, name="Calibri", size=11)
_BOLD_FONT   = Font(bold=True)


def _style_sheet(ws, date_cols=(), num_cols=(), currency_cols=()):
    """Başlık stili + kolon genişliği + sayı formatı uygular."""
    for cell in ws[1]:
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        cell.alignment = Alignment(horizontal="center")

    for col_idx in range(1, ws.max_column + 1):
        col_letter = get_column_letter(col_idx)
        header = ws.cell(1, col_idx).value

        for cell in ws.iter_cols(min_col=col_idx, max_col=col_idx,
                                  min_row=2, max_row=ws.max_row):
            for c in cell:
                if header in date_cols:
                    c.number_format = "YYYY-MM-DD"
                elif header in num_cols:
                    c.number_format = "#,##0.00"
                elif header in currency_cols:
                    c.number_format = "#,##0.00"

        max_len = max(
            len(str(ws.cell(r, col_idx).value or ""))
            for r in range(1, ws.max_row + 1)
        )
        ws.column_dimensions[col_letter].width = min(max_len + 4, 45)


def write_tahmin(df: pd.DataFrame, path: Path) -> None:
    """Tahminlenen_Talep.xlsx dosyasını yazar."""
    eksik = [c for c in TAHMIN_KOLONLAR if c not in df.columns]
    if eksik:
        raise ValueError(f"Tahmin çıktısında eksik kolon: {eksik}")

    df_out = df[TAHMIN_KOLONLAR].copy()
    df_out.to_excel(path, index=False, engine="openpyxl")

    wb = load_workbook(path)
    ws = wb.active
    _style_sheet(ws,
                 date_cols={"Tarih"},
                 num_cols={"Tahmin Edilen Desi"})
    wb.save(path)


def write_plan(df: pd.DataFrame, path: Path, toplam_maliyet: float) -> None:
    """Araç_Planlama.xlsx dosyasını yazar; en sona toplam maliyet satırı ekler."""
    eksik = [c for c in PLAN_KOLONLAR if c not in df.columns]
    if eksik:
        raise ValueError(f"Araç planında eksik kolon: {eksik}")

    df_out = df[PLAN_KOLONLAR].copy()
    df_out.to_excel(path, index=False, engine="openpyxl")

    wb = load_workbook(path)
    ws = wb.active
    _style_sheet(ws,
                 date_cols={"Tarih"},
                 num_cols={"Atanan Desi"},
                 currency_cols={"Maliyet"})

    # Toplam maliyet satırı
    total_row = ws.max_row + 2
    ws.cell(total_row, 1, "TOPLAM MALİYET").font = _BOLD_FONT
    maliyet_col = PLAN_KOLONLAR.index("Maliyet") + 1
    toplam_cell = ws.cell(total_row, maliyet_col, round(toplam_maliyet, 2))
    toplam_cell.font = _BOLD_FONT
    toplam_cell.number_format = "#,##0.00"

    wb.save(path)


def write_unfulfilled(df: pd.DataFrame, path: Path) -> None:
    """Karşılanamayan_Yük.xlsx dosyasını yazar. df boşsa dosya üretilmez."""
    if df.empty:
        return

    eksik = [c for c in UNFUL_KOLONLAR if c not in df.columns]
    if eksik:
        raise ValueError(f"Karşılanamayan tablosunda eksik kolon: {eksik}")

    df_out = df[UNFUL_KOLONLAR].copy()
    df_out.to_excel(path, index=False, engine="openpyxl")

    wb = load_workbook(path)
    ws = wb.active
    _style_sheet(ws,
                 date_cols={"Tarih"},
                 num_cols={"Karşılanamayan Desi"})
    wb.save(path)


def write_all(
    df_tahmin_out: pd.DataFrame,
    tam_plan: pd.DataFrame,
    unfulfilled: pd.DataFrame,
    output_files: dict,
) -> float:
    """
    Tüm çıktı dosyalarını yazar.

    Parametre
    ---------
    df_tahmin_out : rename edilmiş tahmin DataFrame'i (TAHMIN_KOLONLAR)
    tam_plan      : kiralik + ugrama + spot birleşik plan (PLAN_KOLONLAR)
    unfulfilled   : karşılanamayan yük (UNFUL_KOLONLAR); boş olabilir
    output_files  : config.OUTPUT_FILES sözlüğü

    Döndürür
    --------
    float — toplam operasyonel maliyet (TL)
    """
    output_files["tahmin"].parent.mkdir(parents=True, exist_ok=True)

    toplam_maliyet = round(tam_plan["Maliyet"].sum(), 2) if not tam_plan.empty else 0.0

    write_tahmin(df_tahmin_out, output_files["tahmin"])
    write_plan(tam_plan, output_files["plan"], toplam_maliyet)
    write_unfulfilled(unfulfilled, output_files["karsılanamaz"])

    print(f"  Toplam operasyonel maliyet : {toplam_maliyet:>15,.2f} TL")
    print(f"  Tahmin dosyası             : {output_files['tahmin']}")
    print(f"  Araç plan dosyası          : {output_files['plan']}")
    if not unfulfilled.empty:
        print(f"  Karşılanamayan yük         : {output_files['karsılanamaz']}")

    return toplam_maliyet
